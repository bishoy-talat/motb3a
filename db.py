# -*- coding: utf-8 -*-
"""
db.py
طبقة قاعدة البيانات لتطبيق متابعة المخدومين.
تستخدم SQLite (ملف محلي واحد على جهاز المستخدم، بدون إنترنت).
"""

import sqlite3
import os
import shutil
import hashlib
import glob
from datetime import date, datetime, timedelta

# DATA_DIR: افتراضيًا جنب هذا الملف، أو غيّره بمتغير بيئة DATA_DIR (مفيد لو عندك قرص تخزين
# دائم مثبّت في مسار معيّن على الاستضافة، مثل /data على Fly.io)
_DATA_DIR = os.environ.get("DATA_DIR") or os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(_DATA_DIR, "mokhdomeen.db")
PHOTOS_DIR = os.path.join(_DATA_DIR, "photos")
BACKUPS_DIR = os.path.join(_DATA_DIR, "backups")


def get_connection():
    # timeout: تنتظر حتى 15 ثانية لو الملف مقفول مؤقتًا بدل ما تفشل فورًا
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.execute("PRAGMA foreign_keys = ON")
    # WAL يسمح بقراءة/كتابة أكثر سلاسة ويقلل أخطاء "database is locked" على ويندوز
    conn.execute("PRAGMA journal_mode = WAL")
    conn.row_factory = sqlite3.Row
    return conn


def _hash_password(password):
    """
    تجزئة بسيطة لكلمة المرور (SHA-256 + ملح ثابت خاص بالتطبيق).
    ملاحظة: هذا مناسب لأداة محلية صغيرة (خدمة/مدرسة أحد)، وليس بمستوى حماية أنظمة بنكية.
    """
    salt = "mokhdomeen_app_salt_v1"
    return hashlib.sha256((salt + password).encode("utf-8")).hexdigest()


def log_activity(username, action, details=""):
    """يسجّل حدثًا في سجل النشاط (من أضاف/عدّل/حذف إيه ومتى)."""
    conn = get_connection()
    try:
        conn.execute(
            "INSERT INTO activity_log (timestamp, username, action, details) VALUES (?, ?, ?, ?)",
            (datetime.now().strftime("%Y-%m-%d %H:%M:%S"), username or "-", action, details),
        )
        conn.commit()
    finally:
        conn.close()


def get_recent_activity(limit=200):
    conn = get_connection()
    try:
        return conn.execute(
            "SELECT * FROM activity_log ORDER BY id DESC LIMIT ?", (limit,)
        ).fetchall()
    finally:
        conn.close()


def init_db():
    """ينشئ الجداول إذا لم تكن موجودة، ويضيف بيانات ابتدائية بسيطة، ويهاجر قواعد البيانات القديمة."""
    os.makedirs(PHOTOS_DIR, exist_ok=True)
    os.makedirs(BACKUPS_DIR, exist_ok=True)

    conn = get_connection()
    try:
        cur = conn.cursor()

        cur.execute("""
            CREATE TABLE IF NOT EXISTS stages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                meeting_day TEXT
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                phone TEXT,
                birthdate TEXT,        -- بصيغة YYYY-MM-DD
                class_name TEXT,
                stage_id INTEGER,
                active INTEGER DEFAULT 1,
                notes TEXT,
                join_date TEXT,
                photo_path TEXT,
                family_id INTEGER,
                tag TEXT,
                FOREIGN KEY(stage_id) REFERENCES stages(id) ON DELETE SET NULL,
                FOREIGN KEY(family_id) REFERENCES families(id) ON DELETE SET NULL
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS families (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS activity_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                username TEXT,
                action TEXT NOT NULL,
                details TEXT
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                session_date TEXT NOT NULL,   -- YYYY-MM-DD
                present INTEGER NOT NULL,     -- 1 حاضر / 0 غائب
                UNIQUE(student_id, session_date),
                FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS verses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                verse_text TEXT NOT NULL,
                reference TEXT
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'servant',  -- 'admin' أو 'servant'
                stage_id INTEGER,                       -- NULL يعني كل المراحل (للأدمن عادة)
                class_name TEXT,                         -- NULL يعني كل فصول المرحلة (للخادم المسؤول عن مرحلة كاملة)
                FOREIGN KEY(stage_id) REFERENCES stages(id) ON DELETE SET NULL
            )
        """)

        cur.execute("""
            CREATE TABLE IF NOT EXISTS lessons (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                stage_id INTEGER,
                week_date TEXT,        -- تاريخ الأسبوع YYYY-MM-DD (اختياري)
                title TEXT NOT NULL,
                notes TEXT,
                covered INTEGER DEFAULT 0,
                sort_order INTEGER DEFAULT 0,
                FOREIGN KEY(stage_id) REFERENCES stages(id) ON DELETE CASCADE
            )
        """)

        conn.commit()

        # ---------------- ترحيل قواعد بيانات أُنشئت بنسخ أقدم ----------------
        existing_student_cols = [row["name"] for row in cur.execute("PRAGMA table_info(students)").fetchall()]
        if "stage_id" not in existing_student_cols:
            cur.execute("ALTER TABLE students ADD COLUMN stage_id INTEGER REFERENCES stages(id)")
        if "notes" not in existing_student_cols:
            cur.execute("ALTER TABLE students ADD COLUMN notes TEXT")
        if "join_date" not in existing_student_cols:
            cur.execute("ALTER TABLE students ADD COLUMN join_date TEXT")
        if "photo_path" not in existing_student_cols:
            cur.execute("ALTER TABLE students ADD COLUMN photo_path TEXT")
        if "family_id" not in existing_student_cols:
            cur.execute("ALTER TABLE students ADD COLUMN family_id INTEGER REFERENCES families(id)")
        if "tag" not in existing_student_cols:
            cur.execute("ALTER TABLE students ADD COLUMN tag TEXT")

        existing_stage_cols = [row["name"] for row in cur.execute("PRAGMA table_info(stages)").fetchall()]
        if "meeting_day" not in existing_stage_cols:
            cur.execute("ALTER TABLE stages ADD COLUMN meeting_day TEXT")

        existing_user_cols = [row["name"] for row in cur.execute("PRAGMA table_info(users)").fetchall()]
        if "class_name" not in existing_user_cols:
            cur.execute("ALTER TABLE users ADD COLUMN class_name TEXT")
        conn.commit()

        # ترحيل الصور: لو قاعدة البيانات جاية من برنامج الكمبيوتر القديم، مسارات الصور فيها كانت
        # مسارات جهاز تاني (مثلًا C:\\Users\\...). لو نسخت فولدر photos جنب القاعدة، نصلّح المسارات تلقائيًا.
        for row in cur.execute(
            "SELECT id, photo_path FROM students WHERE photo_path IS NOT NULL AND photo_path != ''"
        ).fetchall():
            old_path = row["photo_path"]
            if not os.path.isfile(old_path):
                candidate = os.path.join(PHOTOS_DIR, os.path.basename(old_path.replace("\\", "/")))
                if os.path.isfile(candidate):
                    cur.execute("UPDATE students SET photo_path=? WHERE id=?", (candidate, row["id"]))
        conn.commit()

        # بيانات ابتدائية بسيطة لآيات اليوم (يمكن للمستخدم إضافة المزيد لاحقًا)
        cur.execute("SELECT COUNT(*) AS c FROM verses")
        if cur.fetchone()["c"] == 0:
            seed_verses = [
                ("الرب راعيّ فلا يعوزني شيء.", "مزمور 23:1"),
                ("أحبوا بعضكم بعضًا كما أحببتكم.", "يوحنا 13:34"),
                ("لأن الله أحب العالم حتى بذل ابنه الوحيد.", "يوحنا 3:16"),
                ("كل شيء أستطيع في المسيح الذي يقويني.", "فيلبي 4:13"),
                ("تعالوا إليّ يا جميع المتعبين والثقيلي الأحمال وأنا أريحكم.", "متى 11:28"),
            ]
            cur.executemany(
                "INSERT INTO verses (verse_text, reference) VALUES (?, ?)", seed_verses
            )
            conn.commit()

        # مستخدم أدمن افتراضي لو مفيش مستخدمين خالص (أول تشغيل)
        cur.execute("SELECT COUNT(*) AS c FROM users")
        if cur.fetchone()["c"] == 0:
            cur.execute(
                "INSERT INTO users (username, password_hash, role, stage_id) VALUES (?, ?, 'admin', NULL)",
                # على سيرفر منشور على الإنترنت: اضبط متغير البيئة ADMIN_INITIAL_PASSWORD قبل أول تشغيل
                ("admin", _hash_password(os.environ.get("ADMIN_INITIAL_PASSWORD") or "admin")),
            )
            conn.commit()
    finally:
        conn.close()


# ---------------------- المراحل (Stages) ----------------------

def add_stage(name, meeting_day=None):
    """يضيف مرحلة جديدة. يتجاهل الإضافة بصمت لو الاسم موجود مسبقًا (فريد)."""
    conn = get_connection()
    try:
        conn.execute("INSERT INTO stages (name, meeting_day) VALUES (?, ?)", (name, meeting_day))
        conn.commit()
    except sqlite3.IntegrityError:
        pass  # المرحلة موجودة بالفعل
    finally:
        conn.close()


def get_or_create_stage(name):
    """يرجع id المرحلة، وينشئها لو مش موجودة. مفيد أثناء استيراد إكسل."""
    name = (name or "").strip()
    if not name:
        return None
    conn = get_connection()
    try:
        row = conn.execute("SELECT id FROM stages WHERE name=?", (name,)).fetchone()
        if row:
            return row["id"]
        cur = conn.execute("INSERT INTO stages (name) VALUES (?)", (name,))
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def update_stage(stage_id, new_name, meeting_day=None):
    conn = get_connection()
    try:
        conn.execute("UPDATE stages SET name=?, meeting_day=? WHERE id=?", (new_name, meeting_day, stage_id))
        conn.commit()
    except sqlite3.IntegrityError:
        pass
    finally:
        conn.close()


# نُبقي الاسم القديم كواجهة متوافقة مع كود سابق
def rename_stage(stage_id, new_name):
    update_stage(stage_id, new_name)


def delete_stage(stage_id):
    """يحذف المرحلة. المخدومون التابعون لها لا يُحذفون، فقط تصبح مرحلتهم (بدون مرحلة)."""
    conn = get_connection()
    try:
        conn.execute("UPDATE students SET stage_id=NULL WHERE stage_id=?", (stage_id,))
        conn.execute("DELETE FROM stages WHERE id=?", (stage_id,))
        conn.commit()
    finally:
        conn.close()


def get_all_stages():
    conn = get_connection()
    try:
        return conn.execute("SELECT * FROM stages ORDER BY name").fetchall()
    finally:
        conn.close()


def count_students_in_stage(stage_id):
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT COUNT(*) AS c FROM students WHERE stage_id=? AND active=1", (stage_id,)
        ).fetchone()
        return row["c"]
    finally:
        conn.close()


# ---------------------- العائلات (Families) ----------------------

def get_or_create_family(name):
    name = (name or "").strip()
    if not name:
        return None
    conn = get_connection()
    try:
        row = conn.execute("SELECT id FROM families WHERE name=?", (name,)).fetchone()
        if row:
            return row["id"]
        cur = conn.execute("INSERT INTO families (name) VALUES (?)", (name,))
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def get_all_families():
    conn = get_connection()
    try:
        return conn.execute("SELECT * FROM families ORDER BY name").fetchall()
    finally:
        conn.close()


def get_family_members(family_id, exclude_student_id=None):
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT * FROM students WHERE family_id=? AND active=1 ORDER BY name", (family_id,)
        ).fetchall()
        if exclude_student_id:
            rows = [r for r in rows if r["id"] != exclude_student_id]
        return rows
    finally:
        conn.close()


# ---------------------- المخدومين (Students) ----------------------

TAGS = ["", "يحتاج متابعة خاصة", "منتظم جدًا", "جديد بالخدمة"]

def add_student(name, phone, birthdate, class_name, stage_id=None, notes="", join_date="",
                 photo_path=None, family_id=None, tag=""):
    conn = get_connection()
    try:
        cur = conn.execute(
            """INSERT INTO students (name, phone, birthdate, class_name, stage_id, notes, join_date,
                                       photo_path, family_id, tag)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (name, phone, birthdate, class_name, stage_id, notes, join_date, photo_path, family_id, tag),
        )
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def update_student(student_id, name, phone, birthdate, class_name, stage_id=None, notes="", join_date="",
                     family_id=None, tag=""):
    """تحديث بيانات المخدوم (بدون الصورة - استخدم set_student_photo لتحديث الصورة)."""
    conn = get_connection()
    try:
        conn.execute(
            """UPDATE students SET name=?, phone=?, birthdate=?, class_name=?, stage_id=?, notes=?,
                                     join_date=?, family_id=?, tag=?
               WHERE id=?""",
            (name, phone, birthdate, class_name, stage_id, notes, join_date, family_id, tag, student_id),
        )
        conn.commit()
    finally:
        conn.close()


def set_student_photo(student_id, photo_path):
    conn = get_connection()
    try:
        conn.execute("UPDATE students SET photo_path=? WHERE id=?", (photo_path, student_id))
        conn.commit()
    finally:
        conn.close()


def save_student_photo_file(student_id, source_filepath):
    """
    ينسخ صورة من مكان اختيار المستخدم إلى فولدر photos/ بجانب قاعدة البيانات،
    ويحدّث photo_path في قاعدة البيانات. يرجع المسار الجديد.
    """
    os.makedirs(PHOTOS_DIR, exist_ok=True)
    ext = os.path.splitext(source_filepath)[1].lower() or ".png"
    dest_filename = f"student_{student_id}{ext}"
    dest_path = os.path.join(PHOTOS_DIR, dest_filename)
    shutil.copyfile(source_filepath, dest_path)
    set_student_photo(student_id, dest_path)
    return dest_path


def delete_student(student_id):
    """حذف ناعم (Soft Delete): يضع المخدوم في سلة المحذوفات بدل مسحه نهائيًا."""
    conn = get_connection()
    try:
        conn.execute("UPDATE students SET active=0 WHERE id=?", (student_id,))
        conn.commit()
    finally:
        conn.close()


def restore_student(student_id):
    conn = get_connection()
    try:
        conn.execute("UPDATE students SET active=1 WHERE id=?", (student_id,))
        conn.commit()
    finally:
        conn.close()


def permanently_delete_student(student_id):
    """حذف نهائي لا رجعة فيه - يُستخدم من سلة المحذوفات فقط."""
    conn = get_connection()
    try:
        conn.execute("DELETE FROM students WHERE id=?", (student_id,))
        conn.commit()
    finally:
        conn.close()


def get_deleted_students():
    conn = get_connection()
    try:
        return conn.execute(
            """SELECT s.*, st.name AS stage_name FROM students s
               LEFT JOIN stages st ON st.id = s.stage_id
               WHERE s.active=0 ORDER BY s.name"""
        ).fetchall()
    finally:
        conn.close()


def get_all_students(active_only=True, stage_id=None, search="", class_name=None):
    """
    stage_id:
      None       -> كل المخدومين من كل المراحل
      "unset"    -> فقط المخدومون بدون مرحلة محددة
      رقم (int)  -> فقط مخدومو هذه المرحلة
    class_name:
      None       -> كل الفصول
      "unset"    -> فقط المخدومون بدون فصل محدد
      نص         -> فقط مخدومو هذا الفصل بالضبط
    search: نص للبحث في الاسم (بحث جزئي، غير حساس لحالة الأحرف)
    """
    conn = get_connection()
    try:
        query = """
            SELECT s.*, st.name AS stage_name, f.name AS family_name
            FROM students s
            LEFT JOIN stages st ON st.id = s.stage_id
            LEFT JOIN families f ON f.id = s.family_id
            WHERE 1=1
        """
        params = []
        if active_only:
            query += " AND s.active=1"
        if stage_id == "unset":
            query += " AND s.stage_id IS NULL"
        elif stage_id is not None:
            query += " AND s.stage_id=?"
            params.append(stage_id)
        if class_name == "unset":
            query += " AND (s.class_name IS NULL OR s.class_name='')"
        elif class_name is not None:
            query += " AND s.class_name=?"
            params.append(class_name)
        if search:
            query += " AND s.name LIKE ?"
            params.append(f"%{search}%")
        query += " ORDER BY s.name"
        return conn.execute(query, params).fetchall()
    finally:
        conn.close()


def get_distinct_classes(stage_id=None):
    """
    يرجع قائمة أسماء الفصول الموجودة فعليًا (لمرحلة معينة أو كل المراحل)، الفريدة فقط، مرتبة.
    القيمة الفارغة (بدون فصل) لا تُدرج هنا - تُعامل بشكل منفصل في الواجهة كـ "بدون فصل".
    """
    students = get_all_students(stage_id=stage_id)
    classes = sorted({s["class_name"] for s in students if s["class_name"]})
    return classes


# ---------------------- الحضور والغياب (Attendance) ----------------------

def set_attendance(student_id, session_date, present):
    conn = get_connection()
    try:
        conn.execute(
            """INSERT INTO attendance (student_id, session_date, present)
               VALUES (?, ?, ?)
               ON CONFLICT(student_id, session_date)
               DO UPDATE SET present=excluded.present""",
            (student_id, session_date, 1 if present else 0),
        )
        conn.commit()
    finally:
        conn.close()


def get_attendance_for_date(session_date):
    """يرجع خريطة student_id -> present لتاريخ معيّن."""
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT student_id, present FROM attendance WHERE session_date=?",
            (session_date,),
        ).fetchall()
        return {r["student_id"]: r["present"] for r in rows}
    finally:
        conn.close()


def get_recent_absentees(threshold=2, lookback_sessions=3, stage_id=None):
    """
    يرجع قائمة المخدومين الذين غابوا 'threshold' مرات أو أكثر
    من ضمن آخر 'lookback_sessions' جلسات مسجّلة لنفس المرحلة، لتنبيه الخادم بضرورة الافتقاد.
    تمرير stage_id يضمن أن جلسات مرحلة معينة لا تتداخل مع جلسات مرحلة أخرى
    (فكل مرحلة قد تجتمع في أيام مختلفة).

    ملاحظة: هذه الدالة تحسب على مستوى المرحلة ككل. لعزل كل فصل داخل المرحلة عن الآخر
    (لأن كل فصل قد يجتمع في يوم مختلف عن باقي فصول نفس المرحلة)، استخدم
    get_recent_absentees_grouped بدلًا منها.
    """
    conn = get_connection()
    try:
        stage_filter_sql = ""
        stage_params = []
        if stage_id == "unset":
            stage_filter_sql = " AND s.stage_id IS NULL"
        elif stage_id is not None:
            stage_filter_sql = " AND s.stage_id=?"
            stage_params.append(stage_id)

        session_dates = conn.execute(
            f"""
            SELECT DISTINCT a.session_date
            FROM attendance a
            JOIN students s ON s.id = a.student_id
            WHERE 1=1 {stage_filter_sql}
            ORDER BY a.session_date DESC LIMIT ?
            """,
            (*stage_params, lookback_sessions),
        ).fetchall()
        dates = [r["session_date"] for r in session_dates]
        if not dates:
            return []

        placeholders = ",".join("?" for _ in dates)
        rows = conn.execute(
            f"""
            SELECT s.id, s.name, s.phone, s.class_name,
                   SUM(CASE WHEN a.present=0 THEN 1 ELSE 0 END) AS absences
            FROM students s
            JOIN attendance a ON a.student_id = s.id
            WHERE a.session_date IN ({placeholders}) AND s.active=1 {stage_filter_sql}
            GROUP BY s.id
            HAVING absences >= ?
            ORDER BY absences DESC, s.name
            """,
            (*dates, *stage_params, threshold),
        ).fetchall()
        return rows
    finally:
        conn.close()


def get_recent_absentees_grouped(threshold=2, lookback_sessions=3, stage_id=None, class_name=None):
    """
    مثل get_recent_absentees لكن كل فصل داخل كل مرحلة يُحسب بمعزل تام عن باقي الفصول
    (وكل مرحلة بمعزل عن باقي المراحل) - لأن كل فصل ممكن يجتمع في يوم مختلف تمامًا،
    فجلسات فصل معين لازم متتخلطش مع جلسات فصل تاني حتى لو نفس المرحلة.

    تمرير class_name يضيق النتيجة لفصل واحد بالتحديد (مفيد لخادم مسؤول عن فصل واحد بس).

    يرجع قائمة من dicts بالشكل:
    [{"stage_name": ..., "class_name": ..., "absentees": [صفوف مثل get_recent_absentees]}, ...]
    مرتبة حسب اسم المرحلة ثم اسم الفصل، وبدون أي مجموعة فاضية (بدون غياب مؤهل للتنبيه).
    """
    students = get_all_students(stage_id=stage_id, class_name=class_name)
    if not students:
        return []

    # تجميع المخدومين حسب (المرحلة، الفصل) معًا
    groups = {}
    for s in students:
        key = (s["stage_name"] or "بدون مرحلة", s["class_name"] or "بدون فصل")
        groups.setdefault(key, []).append(s["id"])

    conn = get_connection()
    try:
        results = []
        for (stage_name, class_name) in sorted(groups.keys()):
            member_ids = groups[(stage_name, class_name)]
            placeholders = ",".join("?" for _ in member_ids)

            session_dates = conn.execute(
                f"""SELECT DISTINCT session_date FROM attendance
                    WHERE student_id IN ({placeholders})
                    ORDER BY session_date DESC LIMIT ?""",
                (*member_ids, lookback_sessions),
            ).fetchall()
            dates = [r["session_date"] for r in session_dates]
            if not dates:
                continue

            date_placeholders = ",".join("?" for _ in dates)
            rows = conn.execute(
                f"""
                SELECT s.id, s.name, s.phone,
                       SUM(CASE WHEN a.present=0 THEN 1 ELSE 0 END) AS absences
                FROM students s
                JOIN attendance a ON a.student_id = s.id
                WHERE s.id IN ({placeholders}) AND a.session_date IN ({date_placeholders})
                GROUP BY s.id
                HAVING absences >= ?
                ORDER BY absences DESC, s.name
                """,
                (*member_ids, *dates, threshold),
            ).fetchall()

            if rows:
                results.append({"stage_name": stage_name, "class_name": class_name, "absentees": rows})
        return results
    finally:
        conn.close()


def get_recent_absentees_flat(threshold=2, lookback_sessions=3, stage_id=None, class_name=None):
    """نفس get_recent_absentees_grouped لكن كقائمة واحدة مسطّحة (بدون تجميع) - مفيدة للعرض المختصر."""
    grouped = get_recent_absentees_grouped(threshold=threshold, lookback_sessions=lookback_sessions,
                                             stage_id=stage_id, class_name=class_name)
    flat = []
    for group in grouped:
        flat.extend(group["absentees"])
    return flat


# ---------------------- أعياد الميلاد (Birthdays) ----------------------

def parse_flexible_date(value):
    """
    يحاول تحويل قيمة تاريخ (نص، أو datetime/date) إلى كائن date بايثون،
    ويقبل عدة صيغ شائعة حتى لو لم تُكتب بالصيغة القياسية YYYY-MM-DD:
    YYYY-MM-DD, DD/MM/YYYY, DD-MM-YYYY, MM/DD/YYYY, YYYY/MM/DD.
    يرجع None لو تعذر فهم النص نهائيًا (بدل ما يتجاهله بصمت في مكان بعيد).
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def get_upcoming_birthdays(days_ahead=7, stage_id=None, class_name=None):
    """يرجع المخدومين الذين عيد ميلادهم خلال الأيام القادمة (بصرف النظر عن السنة)."""
    students = get_all_students(stage_id=stage_id, class_name=class_name)
    today = date.today()
    upcoming = []

    for s in students:
        bd = parse_flexible_date(s["birthdate"])
        if bd is None:
            continue

        this_year_bd = _safe_replace_year(bd, today.year)
        if this_year_bd < today:
            this_year_bd = _safe_replace_year(bd, today.year + 1)

        delta = (this_year_bd - today).days
        if 0 <= delta <= days_ahead:
            upcoming.append((s, this_year_bd, delta))

    upcoming.sort(key=lambda x: x[2])
    return upcoming


def _safe_replace_year(d, year):
    """مثل d.replace(year=...) لكن يتعامل مع 29 فبراير في سنة غير كبيسة (يحوّله لـ 28 فبراير)."""
    try:
        return d.replace(year=year)
    except ValueError:
        return d.replace(month=2, day=28, year=year)


# ---------------------- الآية اليومية (Daily Verse) ----------------------

def add_verse(verse_text, reference):
    conn = get_connection()
    try:
        conn.execute(
            "INSERT INTO verses (verse_text, reference) VALUES (?, ?)",
            (verse_text, reference),
        )
        conn.commit()
    finally:
        conn.close()


def get_all_verses():
    conn = get_connection()
    try:
        return conn.execute("SELECT * FROM verses ORDER BY id").fetchall()
    finally:
        conn.close()


def get_verse_of_the_day():
    """يختار آية بشكل ثابت يوميًا (نفس الآية طوال اليوم لكل المستخدمين) بالتدوير حسب يوم السنة."""
    verses = get_all_verses()
    if not verses:
        return None
    day_index = date.today().timetuple().tm_yday
    return verses[day_index % len(verses)]


# ---------------------- استيراد من إكسل (Excel Import) ----------------------

# كلمات مفتاحية للتعرف على كل عمود بصرف النظر عن صياغة العنوان بالضبط
_COLUMN_KEYWORDS = {
    "name": ["اسم", "الاسم", "name"],
    "phone": ["هاتف", "موبايل", "تليفون", "phone", "mobile"],
    "birthdate": ["ميلاد", "تاريخ الميلاد", "birth", "dob"],
    "class_name": ["فصل", "مجموعة", "class", "group"],
    "stage": ["مرحلة", "المرحلة", "stage", "level"],
}


def _match_column(header_text):
    header_text = (header_text or "").strip().lower()
    for field, keywords in _COLUMN_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in header_text:
                return field
    return None


def _parse_excel_date(value):
    """يحاول تحويل قيمة تاريخ من إكسل (قد تكون datetime أو نص) إلى نص YYYY-MM-DD."""
    if value is None or value == "":
        return ""
    parsed = parse_flexible_date(value)
    if parsed:
        return parsed.isoformat()
    return str(value).strip()  # نحتفظ بالنص كما هو لو الصيغة غير معروفة، بدل فقدان البيانات


def import_students_from_excel(filepath):
    """
    يستورد المخدومين من ملف إكسل (.xlsx).
    يتوقع صف عناوين أول (أي ترتيب للأعمدة)، ويتعرف عليها بالكلمات المفتاحية:
      الاسم / الهاتف / تاريخ الميلاد / الفصل / المرحلة
    عمود 'الاسم' إلزامي؛ باقي الأعمدة اختيارية.
    يُنشئ المرحلة تلقائيًا لو مش موجودة.

    ملاحظة تقنية: كل الاستيراد يتم عبر اتصال واحد بقاعدة البيانات (بدل فتح/غلق
    اتصال جديد لكل صف)، لتفادي مشاكل "database is locked" التي قد تحدث على
    ويندوز عند فتح/غلق اتصالات كثيرة ومتتالية بسرعة.

    يرجع dict: {"imported": عدد, "skipped": عدد, "errors": [رسائل]}
    """
    import openpyxl

    result = {"imported": 0, "skipped": 0, "errors": []}

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as e:
        result["errors"].append(f"تعذر فتح الملف: {e}")
        return result

    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        result["errors"].append("الملف فارغ.")
        return result

    header_row = rows[0]
    col_map = {}  # field -> column index
    for idx, header in enumerate(header_row):
        field = _match_column(str(header) if header is not None else "")
        if field and field not in col_map:
            col_map[field] = idx

    if "name" not in col_map:
        result["errors"].append(
            "لم يتم العثور على عمود باسم يحتوي على كلمة 'اسم'. تأكد من وجود عمود للاسم في الصف الأول."
        )
        return result

    conn = get_connection()
    try:
        stage_cache = {}  # اسم المرحلة -> id (لتفادي استعلام مكرر لنفس المرحلة)

        def local_get_or_create_stage(name):
            name = (name or "").strip()
            if not name:
                return None
            if name in stage_cache:
                return stage_cache[name]
            row = conn.execute("SELECT id FROM stages WHERE name=?", (name,)).fetchone()
            if row:
                stage_cache[name] = row["id"]
                return row["id"]
            cur = conn.execute("INSERT INTO stages (name) VALUES (?)", (name,))
            stage_cache[name] = cur.lastrowid
            return cur.lastrowid

        for row_num, row in enumerate(rows[1:], start=2):
            def cell(field):
                idx = col_map.get(field)
                if idx is None or idx >= len(row):
                    return ""
                val = row[idx]
                return "" if val is None else val

            name = str(cell("name")).strip()
            if not name:
                result["skipped"] += 1
                continue

            phone_val = cell("phone")
            phone = "" if phone_val == "" else str(phone_val).strip()

            birthdate = _parse_excel_date(cell("birthdate"))
            class_name = str(cell("class_name")).strip()

            stage_name = str(cell("stage")).strip()
            try:
                stage_id = local_get_or_create_stage(stage_name) if stage_name else None
                conn.execute(
                    "INSERT INTO students (name, phone, birthdate, class_name, stage_id) VALUES (?, ?, ?, ?, ?)",
                    (name, phone, birthdate, class_name, stage_id),
                )
                result["imported"] += 1
            except Exception as e:
                result["skipped"] += 1
                result["errors"].append(f"صف {row_num}: {e}")

        conn.commit()
    finally:
        conn.close()

    return result


# ---------------------- المستخدمون وتسجيل الدخول (Users / Auth) ----------------------

def create_user(username, password, role="servant", stage_id=None, class_name=None):
    """role: 'admin' (كل المراحل وإدارة كاملة) أو 'servant' (مرحلة، أو مرحلة+فصل محدد فقط)."""
    conn = get_connection()
    try:
        conn.execute(
            "INSERT INTO users (username, password_hash, role, stage_id, class_name) VALUES (?, ?, ?, ?, ?)",
            (username.strip(), _hash_password(password), role, stage_id, class_name or None),
        )
        conn.commit()
        return True, ""
    except sqlite3.IntegrityError:
        return False, "اسم المستخدم موجود بالفعل"
    finally:
        conn.close()


def verify_login(username, password):
    """يرجع صف المستخدم لو صحيح، وإلا None."""
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT u.*, st.name AS stage_name FROM users u LEFT JOIN stages st ON st.id = u.stage_id WHERE username=?",
            (username.strip(),),
        ).fetchone()
        if row and row["password_hash"] == _hash_password(password):
            return row
        return None
    finally:
        conn.close()


def get_all_users():
    conn = get_connection()
    try:
        return conn.execute(
            "SELECT u.*, st.name AS stage_name FROM users u LEFT JOIN stages st ON st.id = u.stage_id ORDER BY u.username"
        ).fetchall()
    finally:
        conn.close()


def update_user(user_id, role, stage_id, class_name=None):
    conn = get_connection()
    try:
        conn.execute("UPDATE users SET role=?, stage_id=?, class_name=? WHERE id=?",
                      (role, stage_id, class_name or None, user_id))
        conn.commit()
    finally:
        conn.close()


def reset_user_password(user_id, new_password):
    conn = get_connection()
    try:
        conn.execute("UPDATE users SET password_hash=? WHERE id=?", (_hash_password(new_password), user_id))
        conn.commit()
    finally:
        conn.close()


def delete_user(user_id):
    conn = get_connection()
    try:
        conn.execute("DELETE FROM users WHERE id=?", (user_id,))
        conn.commit()
    finally:
        conn.close()


def count_admins():
    conn = get_connection()
    try:
        row = conn.execute("SELECT COUNT(*) AS c FROM users WHERE role='admin'").fetchone()
        return row["c"]
    finally:
        conn.close()


# ---------------------- نسخ احتياطي (Backup) ----------------------

def backup_database(keep_last=15):
    """
    ينسخ ملف قاعدة البيانات الحالي إلى فولدر backups/ باسم فيه الوقت والتاريخ،
    ويحذف النسخ الزيادة عن 'keep_last' (الأقدم يُحذف أولًا).
    يرجع مسار النسخة الجديدة.
    """
    os.makedirs(BACKUPS_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(BACKUPS_DIR, f"mokhdomeen_backup_{timestamp}.db")
    # نستخدم واجهة النسخ الاحتياطي الرسمية في SQLite (بدل نسخ الملف مباشرة) عشان تطلع نسخة
    # متسقة وكاملة حتى لو فيه تعديلات لسه في ملف WAL ولم تُكتب في الملف الرئيسي بعد.
    src_conn = sqlite3.connect(DB_PATH, timeout=15)
    dst_conn = sqlite3.connect(backup_path)
    try:
        src_conn.backup(dst_conn)
    finally:
        dst_conn.close()
        src_conn.close()

    existing = sorted(glob.glob(os.path.join(BACKUPS_DIR, "mokhdomeen_backup_*.db")))
    if len(existing) > keep_last:
        for old_file in existing[: len(existing) - keep_last]:
            try:
                os.remove(old_file)
            except OSError:
                pass

    return backup_path


def list_backups():
    files = sorted(glob.glob(os.path.join(BACKUPS_DIR, "mokhdomeen_backup_*.db")), reverse=True)
    return files


# ---------------------- الخطة السنوية (Lessons) ----------------------

def add_lesson(stage_id, title, week_date="", notes=""):
    conn = get_connection()
    try:
        row = conn.execute("SELECT COALESCE(MAX(sort_order), 0) AS m FROM lessons WHERE stage_id=?",
                            (stage_id,)).fetchone()
        next_order = (row["m"] or 0) + 1
        conn.execute(
            "INSERT INTO lessons (stage_id, title, week_date, notes, sort_order) VALUES (?, ?, ?, ?, ?)",
            (stage_id, title, week_date, notes, next_order),
        )
        conn.commit()
    finally:
        conn.close()


def get_lessons(stage_id=None):
    conn = get_connection()
    try:
        if stage_id is not None:
            return conn.execute(
                "SELECT * FROM lessons WHERE stage_id=? ORDER BY sort_order, id", (stage_id,)
            ).fetchall()
        return conn.execute("SELECT * FROM lessons ORDER BY sort_order, id").fetchall()
    finally:
        conn.close()


def toggle_lesson_covered(lesson_id, covered):
    conn = get_connection()
    try:
        conn.execute("UPDATE lessons SET covered=? WHERE id=?", (1 if covered else 0, lesson_id))
        conn.commit()
    finally:
        conn.close()


def delete_lesson(lesson_id):
    conn = get_connection()
    try:
        conn.execute("DELETE FROM lessons WHERE id=?", (lesson_id,))
        conn.commit()
    finally:
        conn.close()


# ---------------------- إحصائيات الحضور (Attendance Stats) ----------------------

def get_attendance_stats(stage_id=None, class_name=None, min_sessions=1):
    """
    يرجع لكل مخدوم: عدد الجلسات المسجّلة له، عدد مرات الحضور، ونسبة الحضور المئوية.
    مرتبة تنازليًا حسب النسبة (الأكثر انتظامًا أولًا).
    min_sessions: أقل عدد جلسات مسجّلة حتى تُحتسب النسبة (لتفادي 100% وهمية لمخدوم له جلسة واحدة فقط).
    """
    conn = get_connection()
    try:
        stage_filter_sql = ""
        params = []
        if stage_id == "unset":
            stage_filter_sql = " AND s.stage_id IS NULL"
        elif stage_id is not None:
            stage_filter_sql = " AND s.stage_id=?"
            params.append(stage_id)

        class_filter_sql = ""
        if class_name == "unset":
            class_filter_sql = " AND (s.class_name IS NULL OR s.class_name='')"
        elif class_name is not None:
            class_filter_sql = " AND s.class_name=?"
            params.append(class_name)

        rows = conn.execute(
            f"""
            SELECT s.id, s.name, s.phone, st.name AS stage_name,
                   COUNT(a.id) AS total_sessions,
                   SUM(CASE WHEN a.present=1 THEN 1 ELSE 0 END) AS present_count
            FROM students s
            LEFT JOIN attendance a ON a.student_id = s.id
            LEFT JOIN stages st ON st.id = s.stage_id
            WHERE s.active=1 {stage_filter_sql} {class_filter_sql}
            GROUP BY s.id
            HAVING total_sessions >= ?
            ORDER BY (CAST(present_count AS FLOAT) / total_sessions) DESC, s.name
            """,
            (*params, min_sessions),
        ).fetchall()

        stats = []
        for r in rows:
            pct = round((r["present_count"] / r["total_sessions"]) * 100) if r["total_sessions"] else 0
            stats.append({
                "id": r["id"], "name": r["name"], "phone": r["phone"], "stage_name": r["stage_name"],
                "total_sessions": r["total_sessions"], "present_count": r["present_count"], "percentage": pct,
            })
        return stats
    finally:
        conn.close()


def get_all_session_dates(stage_id=None, class_name=None):
    """يرجع كل تواريخ الجلسات المسجّلة لمرحلة (وفصل) معينة (أو الكل)، الأقدم أولًا."""
    conn = get_connection()
    try:
        stage_filter_sql = ""
        params = []
        if stage_id == "unset":
            stage_filter_sql = " AND s.stage_id IS NULL"
        elif stage_id is not None:
            stage_filter_sql = " AND s.stage_id=?"
            params.append(stage_id)
        if class_name == "unset":
            stage_filter_sql += " AND (s.class_name IS NULL OR s.class_name='')"
        elif class_name is not None:
            stage_filter_sql += " AND s.class_name=?"
            params.append(class_name)

        rows = conn.execute(
            f"""
            SELECT DISTINCT a.session_date
            FROM attendance a
            JOIN students s ON s.id = a.student_id
            WHERE 1=1 {stage_filter_sql}
            ORDER BY a.session_date
            """,
            params,
        ).fetchall()
        return [r["session_date"] for r in rows]
    finally:
        conn.close()


# ---------------------- تصدير تقارير إلى إكسل (Excel Reports) ----------------------

def export_attendance_report(filepath, stage_id=None, class_name=None):
    """
    يُصدّر تقرير حضور كامل لمرحلة معينة (أو كل المراحل): صف لكل مخدوم،
    عمود لكل تاريخ جلسة (✓ / ✗ / -)، وعمود أخير بنسبة الحضور الكلية.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    session_dates = get_all_session_dates(stage_id=stage_id, class_name=class_name)
    students = get_all_students(stage_id=stage_id, class_name=class_name)

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "تقرير الحضور"

    header = ["الاسم", "الهاتف", "المرحلة"] + session_dates + ["نسبة الحضور"]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F5DFF", end_color="4F5DFF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    conn = get_connection()
    try:
        for s in students:
            att_rows = conn.execute(
                "SELECT session_date, present FROM attendance WHERE student_id=?", (s["id"],)
            ).fetchall()
            att_map = {r["session_date"]: r["present"] for r in att_rows}

            row = [s["name"], s["phone"] or "", s["stage_name"] or "بدون مرحلة"]
            total, present = 0, 0
            for d in session_dates:
                if d in att_map:
                    total += 1
                    if att_map[d]:
                        present += 1
                        row.append("✓")
                    else:
                        row.append("✗")
                else:
                    row.append("-")
            pct = round((present / total) * 100) if total else 0
            row.append(f"{pct}%")
            sheet.append(row)
    finally:
        conn.close()

    for col_cells in sheet.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        sheet.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 8), 20)

    wb.save(filepath)
    return filepath


def export_phone_list(filepath, stage_id=None, class_name=None):
    """يُصدّر قائمة بأسماء وأرقام هواتف المخدومين (لمرحلة معينة أو الكل) إلى إكسل."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    students = get_all_students(stage_id=stage_id, class_name=class_name)

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "أرقام الهواتف"

    header = ["الاسم", "الهاتف", "المرحلة", "الفصل"]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F5DFF", end_color="4F5DFF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for s in students:
        sheet.append([s["name"], s["phone"] or "", s["stage_name"] or "بدون مرحلة", s["class_name"] or ""])

    for col_cells in sheet.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        sheet.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 8), 25)

    wb.save(filepath)
    return filepath


# ---------------------- التحقق من رقم الهاتف المصري (Phone Validation) ----------------------

def is_valid_egyptian_phone(phone):
    """
    يتحقق أن الرقم مصري محلي صحيح: 11 رقمًا بالضبط، يبدأ بـ 010 أو 011 أو 012 أو 015.
    نص فاضي يعتبر صحيحًا (الهاتف اختياري).
    """
    phone = (phone or "").strip()
    if not phone:
        return True
    digits = "".join(ch for ch in phone if ch.isdigit())
    if len(digits) != 11:
        return False
    return digits[:3] in ("010", "011", "012", "015")


# ---------------------- حساب العمر (Age) ----------------------

def calculate_age(birthdate_text):
    """يرجع العمر بالسنوات من نص تاريخ ميلاد (يقبل نفس الصيغ المرنة)، أو None لو تعذر الفهم."""
    bd = parse_flexible_date(birthdate_text)
    if bd is None:
        return None
    today = date.today()
    years = today.year - bd.year - ((today.month, today.day) < (bd.month, bd.day))
    return years


# ---------------------- لوحة المعلومات (Dashboard) ----------------------

def get_dashboard_stats(stage_id=None, class_name=None):
    """يجمع أهم الأرقام لعرضها في الصفحة الرئيسية بعد تسجيل الدخول."""
    students = get_all_students(stage_id=stage_id, class_name=class_name)
    total_students = len(students)

    today_str = date.today().isoformat()
    today_attendance = get_attendance_for_date(today_str)
    relevant_ids = {s["id"] for s in students}
    present_today = sum(1 for sid, present in today_attendance.items() if present and sid in relevant_ids)
    recorded_today = sum(1 for sid in today_attendance if sid in relevant_ids)
    attendance_pct_today = round((present_today / recorded_today) * 100) if recorded_today else None

    absentees = get_recent_absentees_flat(threshold=2, lookback_sessions=3, stage_id=stage_id, class_name=class_name)
    upcoming_birthdays = get_upcoming_birthdays(days_ahead=7, stage_id=stage_id, class_name=class_name)

    return {
        "total_students": total_students,
        "present_today": present_today,
        "recorded_today": recorded_today,
        "attendance_pct_today": attendance_pct_today,
        "needs_visit_count": len(absentees),
        "needs_visit": absentees[:5],
        "nearest_birthday": upcoming_birthdays[0] if upcoming_birthdays else None,
        "upcoming_birthdays_count": len(upcoming_birthdays),
    }


def get_stages_missing_attendance_today():
    """يرجع المراحل التي معادها اليوم (حسب يوم الأسبوع) ولسه محدش سجّل حضور فيها النهاردة."""
    weekday_names = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
    today_name = weekday_names[date.today().weekday()]
    today_str = date.today().isoformat()

    stages = get_all_stages()
    result = []
    for s in stages:
        if s["meeting_day"] != today_name:
            continue
        conn = get_connection()
        try:
            row = conn.execute(
                """SELECT COUNT(*) AS c FROM attendance a
                   JOIN students st ON st.id = a.student_id
                   WHERE st.stage_id=? AND a.session_date=?""",
                (s["id"], today_str),
            ).fetchone()
        finally:
            conn.close()
        if row["c"] == 0:
            result.append(s)
    return result


# ---------------------- سجل حضور فردي لكل مخدوم (Student Attendance History) ----------------------

def get_student_attendance_history(student_id):
    """يرجع كل سجلات حضور/غياب مخدوم معيّن، الأحدث أولًا."""
    conn = get_connection()
    try:
        return conn.execute(
            "SELECT session_date, present FROM attendance WHERE student_id=? ORDER BY session_date DESC",
            (student_id,),
        ).fetchall()
    finally:
        conn.close()


# ---------------------- استرجاع نسخة احتياطية (Restore Backup) ----------------------

def restore_from_backup(backup_path):
    """
    يستبدل قاعدة البيانات الحالية بنسخة احتياطية محفوظة.
    يأخذ نسخة أمان من الحالة الحالية أولًا قبل الاستبدال (احتياطًا لو حصل خطأ).
    يتطلب إعادة تشغيل البرنامج بعدها.
    """
    if os.path.exists(DB_PATH):
        safety_copy = DB_PATH + ".before_restore"
        shutil.copyfile(DB_PATH, safety_copy)

    shutil.copyfile(backup_path, DB_PATH)

    # حذف ملفات WAL/SHM القديمة حتى لا تتعارض مع الملف المستبدَل
    for suffix in ("-wal", "-shm"):
        stale = DB_PATH + suffix
        if os.path.exists(stale):
            try:
                os.remove(stale)
            except OSError:
                pass


# ---------------------- تصدير PDF (بدعم عربي صحيح الاتجاه) ----------------------

def _shape_arabic(text):
    """يهيئ نص عربي للعرض الصحيح في PDF (ربط الحروف + ترتيب من اليمين لليسار)."""
    import arabic_reshaper
    from bidi.algorithm import get_display
    return get_display(arabic_reshaper.reshape(str(text)))


def _register_arabic_font():
    """يسجّل خطًا يدعم العربية لـ reportlab. يرجع اسم الخط المسجَّل."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts", "DejaVuSans.ttf"),  # خط مرفق مع السيرفر
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
        "C:\\Windows\\Fonts\\arial.ttf",
        "C:\\Windows\\Fonts\\tahoma.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("AppArabicFont", path))
                return "AppArabicFont"
            except Exception:
                continue
    return "Helvetica"  # fallback أخير (لن يعرض العربية بشكل صحيح لكن يمنع تعطل البرنامج)


def export_lessons_pdf(filepath, stage_id, stage_name):
    """يُصدّر الخطة السنوية لمرحلة معينة كملف PDF جاهز للطباعة."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.lib.units import cm

    font_name = _register_arabic_font()
    lessons = get_lessons(stage_id=stage_id)

    c = pdf_canvas.Canvas(filepath, pagesize=A4)
    width, height = A4
    y = height - 2 * cm

    c.setFont(font_name, 16)
    c.drawRightString(width - 2 * cm, y, _shape_arabic(f"الخطة السنوية - {stage_name}"))
    y -= 1.2 * cm

    c.setFont(font_name, 11)
    for lesson in lessons:
        if y < 2 * cm:
            c.showPage()
            c.setFont(font_name, 11)
            y = height - 2 * cm
        status = "✓ تم" if lesson["covered"] else "لم يُغطَّ بعد"
        line = f'{lesson["title"]}   —   {lesson["week_date"] or ""}   ({status})'
        c.drawRightString(width - 2 * cm, y, _shape_arabic(line))
        y -= 0.8 * cm

    if not lessons:
        c.drawRightString(width - 2 * cm, y, _shape_arabic("لا توجد دروس مضافة لهذه المرحلة بعد."))

    c.save()
    return filepath


def export_stats_pdf(filepath, stage_id=None, stage_label="كل المراحل", class_name=None):
    """يُصدّر إحصائيات نسب الحضور كملف PDF جاهز للطباعة."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.lib.units import cm

    font_name = _register_arabic_font()
    stats = get_attendance_stats(stage_id=stage_id, class_name=class_name, min_sessions=1)

    c = pdf_canvas.Canvas(filepath, pagesize=A4)
    width, height = A4
    y = height - 2 * cm

    c.setFont(font_name, 16)
    c.drawRightString(width - 2 * cm, y, _shape_arabic(f"إحصائيات الحضور - {stage_label}"))
    y -= 1.2 * cm

    c.setFont(font_name, 11)
    for s in stats:
        if y < 2 * cm:
            c.showPage()
            c.setFont(font_name, 11)
            y = height - 2 * cm
        line = f'{s["name"]}   —   {s["percentage"]}%   ({s["present_count"]}/{s["total_sessions"]})'
        c.drawRightString(width - 2 * cm, y, _shape_arabic(line))
        y -= 0.8 * cm

    if not stats:
        c.drawRightString(width - 2 * cm, y, _shape_arabic("لا توجد بيانات حضور كافية بعد."))

    c.save()
    return filepath


