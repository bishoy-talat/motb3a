# -*- coding: utf-8 -*-
import os, sys, time, subprocess, tempfile
from datetime import datetime, timedelta, timezone

BASE = "http://127.0.0.1:8890"
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(HERE, "desktop"))
SHOTS = "/tmp/e2e_shots"
os.makedirs(SHOTS, exist_ok=True)

import requests
import tkinter as tk
from tkinter import messagebox, filedialog, simpledialog

# --- نمنع نوافذ الحوار المعطِّلة ونلتقط رسائلها ---
MESSAGES = []
for _n in ("showinfo", "showwarning", "showerror"):
    setattr(messagebox, _n, lambda title, msg, _n=_n, **kw: MESSAGES.append((_n, msg)))
messagebox.askyesno = lambda *a, **k: True
simpledialog.askstring = lambda *a, **k: "عائلة الاختبار"
NEXT_SAVE_PATH = {"path": None}
filedialog.asksaveasfilename = lambda **k: NEXT_SAVE_PATH["path"]

import api_client as client
client.save_config(BASE)
import main as desktop

RESULTS = []
def check(label, cond, extra=""):
    RESULTS.append((label, bool(cond)))
    print(("PASS " if cond else "FAIL ") + label + (f"  [{extra}]" if (extra and not cond) else ""))

def shot(app, name):
    app.update_idletasks(); app.update()
    subprocess.run(["import", "-window", "root", f"{SHOTS}/{name}.png"], capture_output=True)

# ---------- "تطبيق الموبايل": عميل خام بيكلّم الـ API مباشرة (زي تطبيق الويب/Flutter) ----------
class Mobile:
    def __init__(self, username, password):
        r = requests.post(BASE + "/api/login", json={"username": username, "password": password})
        r.raise_for_status()
        self.h = {"Authorization": "Bearer " + r.json()["token"]}
    def get(self, path, **params):
        return requests.get(BASE + path, headers=self.h, params=params)
    def post(self, path, body=None, **kw):
        return requests.post(BASE + path, headers=self.h, json=body, **kw)
    def put(self, path, body):
        return requests.put(BASE + path, headers=self.h, json=body)
    def delete(self, path):
        return requests.delete(BASE + path, headers=self.h)

# =====================================================================
print("=== A) البيانات اتعملت من الموبايل ===")
mob = Mobile("admin", "admin")
mob.post("/api/stages", {"name": "ثانوي", "meeting_day": "الجمعة"})
stage = mob.get("/api/stages").json()[0]
SID = stage["id"]
for name, phone, cls in [("أحمد من الموبايل", "01000000001", "فصل أ"),
                         ("محمد من الموبايل", "01000000002", "فصل ب"),
                         ("سارة من الموبايل", "01000000003", "فصل ب")]:
    r = mob.post("/api/students", {"name": name, "phone": phone, "class_name": cls, "stage_id": SID})
    assert r.status_code == 200, r.text
ids = {s["name"]: s["id"] for s in mob.get("/api/students").json()}
check("الموبايل أنشأ 3 مخدومين", len(ids) == 3)
r = mob.post("/api/users", {"username": "khadem_b", "password": "pass123", "role": "servant",
                            "stage_id": SID, "class_name": "فصل ب"})
check("الموبايل (أدمن) أنشأ خادم مقفول على فصل ب", r.status_code == 200, r.text)

# =====================================================================
print("=== B) برنامج الكمبيوتر بيشوف بيانات الموبايل فورًا ===")
user = client.verify_login("admin", "admin")
check("تسجيل دخول الكمبيوتر عبر السيرفر", user and user["role"] == "admin")
app = desktop.MokhdomeenApp(current_user=user)
app.show_page("students"); app.update()
sp = app.pages["students"]
names_in_table = {sp.tree.item(i, "values")[1] for i in sp.tree.get_children()}
check("جدول المخدومين في الكمبيوتر فيه مخدومين الموبايل", {"أحمد من الموبايل", "محمد من الموبايل", "سارة من الموبايل"} <= names_in_table, str(names_in_table))
shot(app, "01_students_desktop")

print("=== C) إضافة من الكمبيوتر تظهر في الموبايل ===")
sp.name_var.set("مينا من الكمبيوتر"); sp.phone_var.set("01111111111")
sp.class_var.set("فصل أ"); sp.stage_var.set("ثانوي"); sp.birth_var.set("30/08/2000")
sp.add_student(); app.update()
mobile_list = [s["name"] for s in mob.get("/api/students").json()]
check("المخدوم المضاف من الكمبيوتر ظهر للموبايل", "مينا من الكمبيوتر" in mobile_list)
ids = {s["name"]: s["id"] for s in mob.get("/api/students").json()}
bd = [s for s in mob.get("/api/students").json() if s["name"] == "مينا من الكمبيوتر"][0]["birthdate"]
check("تاريخ الميلاد اتوحّد لصيغة ISO", bd == "2000-08-30", bd)

# قواعد التحقق موحّدة: الكمبيوتر بيرفض قبل ما يوصل السيرفر، والسيرفر بيرفض لو جه من الموبايل
r = mob.post("/api/students", {"name": "اسم123", "stage_id": SID})
check("السيرفر يرفض اسم فيه أرقام (للموبايل)", r.status_code == 400)
r = mob.post("/api/students", {"name": "فصل غلط", "class_name": "فصل1", "stage_id": SID})
check("السيرفر يرفض فصل فيه أرقام", r.status_code == 400)
r = mob.post("/api/students", {"name": "تليفون غلط", "phone": "123", "stage_id": SID})
check("السيرفر يرفض رقم هاتف غير صحيح", r.status_code == 400)

print("=== D) الحضور: كمبيوتر → موبايل → كمبيوتر ===")
app.show_page("attendance"); app.update()
ap = app.pages["attendance"]
check("صفحة الحضور فيها كل المخدومين", len(ap.check_vars) == 4, str(len(ap.check_vars)))
ap.check_vars[ids["أحمد من الموبايل"]] = True
ap.check_vars[ids["مينا من الكمبيوتر"]] = True
ap.save_attendance(); app.update()
today = datetime.now().date().isoformat()
m = mob.get("/api/attendance/map").json()
check("الحضور المحفوظ من الكمبيوتر ظاهر للموبايل (كقيم منطقية حقيقية true/false)", m.get(str(ids["أحمد من الموبايل"])) is True and m.get(str(ids["مينا من الكمبيوتر"])) is True and m.get(str(ids["محمد من الموبايل"])) is False, str(m))
td = mob.get("/api/attendance/today").json()["existing"]
check("/attendance/today بترجع true/false حقيقية (مهم لتطبيق Flutter)", all(isinstance(v, bool) for v in td.values()) and td, str(td))
mob.post("/api/attendance", {"records": {str(ids["سارة من الموبايل"]): True}})
ap.refresh(); app.update()
check("حضور سجّله الموبايل ظهر في جدول الكمبيوتر", ap.check_vars.get(ids["سارة من الموبايل"]) is True)
shot(app, "02_attendance_desktop")

print("=== E) عزل الفصول والغياب المتكرر ===")
for d in ("2026-08-01", "2026-08-08", "2026-08-15"):
    mob.post("/api/attendance", {"date": d, "records": {str(ids["محمد من الموبايل"]): False}})
groups = client.get_recent_absentees_grouped()
flat = {a["name"]: g["class_name"] for g in groups for a in g["absentees"]}
check("محمد (فصل ب) ظهر في قائمة الافتقاد", flat.get("محمد من الموبايل") == "فصل ب", str(flat))
ap.refresh(); app.update()
check("قائمة الافتقاد في واجهة الكمبيوتر ظهر فيها عنوان فصل ب", any("فصل ب" in w.cget("text") for w in ap.alert_container.winfo_children() if isinstance(w, tk.Label)))

print("=== F) لوحة المعلومات وأعياد الميلاد ===")
app.show_page("dashboard"); app.update()
st = client.get_dashboard_stats()
check("إحصائيات الرئيسية بترجع أرقام صحيحة", st["total_students"] == 4 and st["needs_visit_count"] >= 1, str(st)[:200])
mob.put(f"/api/students/{ids['سارة من الموبايل']}", {"name": "سارة من الموبايل", "phone": "01000000003",
        "class_name": "فصل ب", "stage_id": SID, "birthdate": datetime.now().strftime("1999-%m-%d")})
b = client.get_upcoming_birthdays()
check("عيد ميلاد اليوم طلع بصيغة (مخدوم, date, delta)", len(b) == 1 and b[0][2] == 0 and hasattr(b[0][1], "strftime"), str(b))
app.show_page("birthdays"); app.update()
shot(app, "03_dashboard_birthdays")

print("=== G) التقارير والتصدير عبر واجهة الكمبيوتر ===")
app.show_page("reports"); app.update()
rp = app.pages["reports"]
import openpyxl
NEXT_SAVE_PATH["path"] = "/tmp/e2e_att.xlsx"; rp.export_attendance()
wb = openpyxl.load_workbook("/tmp/e2e_att.xlsx")
check("تقرير الحضور (Excel) اتنزّل وبيتفتح وفيه صفوف", wb.active.max_row >= 4, str(wb.active.max_row))
NEXT_SAVE_PATH["path"] = "/tmp/e2e_phones.xlsx"; rp.export_phones()
check("قائمة الهواتف (Excel) اتنزّلت", openpyxl.load_workbook("/tmp/e2e_phones.xlsx").active.max_row == 5)
NEXT_SAVE_PATH["path"] = "/tmp/e2e_stats.pdf"; rp.export_stats_pdf()
check("إحصائيات PDF اتنزّلت", open("/tmp/e2e_stats.pdf", "rb").read(4) == b"%PDF")
mob.post("/api/lessons", {"stage_id": SID, "title": "الدرس الأول", "week_date": "2026-10-10"})
app.show_page("lessons"); app.update()
lp = app.pages["lessons"]
NEXT_SAVE_PATH["path"] = "/tmp/e2e_lessons.pdf"; lp.export_pdf()
check("الخطة السنوية PDF اتنزّلت", os.path.exists("/tmp/e2e_lessons.pdf") and open("/tmp/e2e_lessons.pdf", "rb").read(4) == b"%PDF")
shot(app, "04_reports_desktop")

print("=== H) الصور واستيراد إكسل ===")
from PIL import Image
Image.new("RGB", (80, 80), (79, 93, 255)).save("/tmp/e2e_photo.png")
client.save_student_photo_file(ids["أحمد من الموبايل"], "/tmp/e2e_photo.png")
stu = mob.get(f"/api/students/{ids['أحمد من الموبايل']}").json()
check("السيرفر مبيسرّبش مسار الملف الحقيقي (اسم الملف بس)", stu["photo_path"] and "/" not in stu["photo_path"], str(stu["photo_path"]))
local = client.fetch_student_photo(ids["أحمد من الموبايل"])
check("الكمبيوتر نزّل الصورة من السيرفر وبتتفتح", local and Image.open(local).size == (80, 80))
check("مخدوم بدون صورة بيرجّع None", client.fetch_student_photo(ids["سارة من الموبايل"]) is None)
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["اسم", "هاتف", "مرحلة", "فصل"]); ws.append(["مستورد من إكسل", "01222222222", "ثانوي", "فصل ج"])
wb.save("/tmp/e2e_import.xlsx")
res = client.import_students_from_excel("/tmp/e2e_import.xlsx")
check("استيراد إكسل عبر السيرفر", res["imported"] == 1, str(res))

print("=== I) المستخدمون والأمان ===")
ok, msg = client.create_user("khadem_b", "x1234", "servant", SID, None)
check("اسم مستخدم مكرر بيرجع (False, رسالة)", ok is False and msg)
users = client.get_all_users()
check("قائمة المستخدمين مفيهاش password_hash", all("password_hash" not in u for u in users))
admin_id = [u for u in users if u["username"] == "admin"][0]["id"]
for label, fn in (("حذف آخر أدمن/الحساب الحالي بيتمنع", lambda: client.delete_user(admin_id)),):
    try:
        fn(); check(label, False)
    except client.ApiError as e:
        check(label, e.status == 400, e.message)
sp.refresh(); app.update()
r = requests.get(BASE + "/api/users", headers={"Authorization": "Bearer bad"})
check("توكن غلط = 401", r.status_code == 401)

print("=== J) النسخ الاحتياطي (على السيرفر) ===")
name = client.backup_database()
check("نسخة احتياطية يدوية اتعملت", name.endswith(".db") and name in client.list_backups())
mob.post("/api/students", {"name": "هيتمسح بعد الاسترجاع", "stage_id": SID})
n_before = len(mob.get("/api/students").json())
client.restore_from_backup(name)
try:
    client.get_all_stages(); check("الجلسات اتلغت بعد الاسترجاع", False)
except client.SessionExpired:
    check("الجلسات اتلغت بعد الاسترجاع (لازم دخول من جديد)", True)
client.verify_login("admin", "admin")
mob = Mobile("admin", "admin")
n_after = len(mob.get("/api/students").json())
check("الاسترجاع رجّع البيانات لوقت النسخة", n_after == n_before - 1, f"{n_before}->{n_after}")
r = mob.post("/api/backup/restore", {"name": "../../etc/passwd"})
check("منع Path Traversal في الاسترجاع", r.status_code == 400)

app.destroy()

# =====================================================================
print("=== K) خادم مقفول على فصل ب يستخدم برنامج الكمبيوتر ===")
suser = client.verify_login("khadem_b", "pass123")
check("الخادم سجّل دخول من الكمبيوتر", suser and suser["class_name"] == "فصل ب")
sapp = desktop.MokhdomeenApp(current_user=suser)
sapp.show_page("students"); sapp.update()
ssp = sapp.pages["students"]
snames = {ssp.tree.item(i, "values")[1] for i in ssp.tree.get_children()}
check("جدول الخادم فيه مخدومين فصل ب بس", snames == {"محمد من الموبايل", "سارة من الموبايل"}, str(snames))
sapp.show_page("attendance"); sapp.update()
sap = sapp.pages["attendance"]
check("حضور الخادم فيه فصله بس", set(sap.check_vars) == {ids["محمد من الموبايل"], ids["سارة من الموبايل"]}, str(sap.check_vars))
shot(sapp, "05_servant_attendance")
check("الخادم مش شايف صفحات الأدمن", "users" not in sapp.pages and "stages" not in sapp.pages)
sapp.destroy()

smob = Mobile("khadem_b", "pass123")
r = smob.get("/api/students", class_name="فصل أ")
check("الخادم بيحاول يطلب فصل تاني: السيرفر يتجاهل", {s["name"] for s in r.json()} == {"محمد من الموبايل", "سارة من الموبايل"})
check("تعديل مخدوم خارج النطاق = 404", smob.put(f"/api/students/{ids['أحمد من الموبايل']}", {"name": "اختراق", "stage_id": SID}).status_code == 404)
check("حذف مخدوم خارج النطاق = 404", smob.delete(f"/api/students/{ids['أحمد من الموبايل']}").status_code == 404)
check("حضور لمخدوم خارج النطاق = 403", smob.post("/api/attendance", {"records": {str(ids["أحمد من الموبايل"]): True}}).status_code == 403)
check("تاريخ حضور مخدوم خارج النطاق = 404", smob.get(f"/api/students/{ids['أحمد من الموبايل']}/attendance_history").status_code == 404)
check("صورة مخدوم خارج النطاق = 404", smob.get(f"/api/students/{ids['أحمد من الموبايل']}/photo").status_code == 404)
check("تصدير الخادم مقصور على فصله", len(openpyxl.load_workbook(__import__("io").BytesIO(smob.get("/api/export/phone_list").content)).active["A"]) == 3)
check("الخادم مينفعش يستورد إكسل (أدمن فقط)", smob.post("/api/students/import_excel", files={"file": ("x.xlsx", b"x")}).status_code == 403)
check("الخادم مينفعش يفتح المستخدمين/السجل", smob.get("/api/users").status_code == 403 and smob.get("/api/activity_log").status_code == 403)
check("الخادم مينفعش يعمل نسخة احتياطية", smob.post("/api/backup").status_code == 403)
# خادم يعدّل على طالب بنفس فصله: الفصل بيتثبّت تلقائيًا
r = smob.put(f"/api/students/{ids['محمد من الموبايل']}", {"name": "محمد معدّل", "class_name": "فصل أ", "stage_id": 999})
fixed = [s for s in smob.get("/api/students").json() if s["id"] == ids["محمد من الموبايل"]][0]
check("الخادم مقدرش ينقل مخدوم لفصل/مرحلة تانية", r.status_code == 200 and fixed["class_name"] == "فصل ب" and fixed["stage_id"] == SID, str(fixed))

print("=== L) سلوك الشبكة والأمان العام ===")
old = client.SERVER_URL
client.SERVER_URL = "http://127.0.0.1:9"
try:
    client.get_all_stages(); check("فشل الاتصال بيرجع رسالة عربية", False)
except client.ApiError as e:
    check("فشل الاتصال بيرجع رسالة عربية واضحة", "تعذر الاتصال" in e.message, e.message)
client.SERVER_URL = old
for i in range(5):
    requests.post(BASE + "/api/login", json={"username": "admin", "password": "wrong"})
try:
    client.verify_login("admin", "admin"); check("قفل الحساب بعد 5 محاولات غلط", False)
except client.ApiError as e:
    check("قفل الحساب بعد 5 محاولات غلط (429)", e.status == 429, e.message)
h = requests.get(BASE + "/api/health").json()["time"]
server_now = datetime.fromisoformat(h)
offset_h = round((server_now - datetime.now(timezone.utc).replace(tzinfo=None)).total_seconds() / 3600)
check("السيرفر شغال بتوقيت مصر (UTC+2 أو +3)", offset_h in (2, 3), f"offset={offset_h}")
r = requests.get(BASE + "/..%2fserver%2fdb.py")
check("منع قراءة ملفات السيرفر عبر مسار الويب", "def get_connection" not in r.text)
check("تطبيق الويب لسه بيتقدم من نفس السيرفر", requests.get(BASE + "/").status_code == 200 and "manifest" in requests.get(BASE + "/").text)

# =====================================================================
passed = sum(1 for _, ok in RESULTS if ok)
print(f"\n===== النتيجة: {passed}/{len(RESULTS)} نجح =====")
for label, ok in RESULTS:
    if not ok:
        print("  فشل:", label)
sys.exit(0 if passed == len(RESULTS) else 1)
